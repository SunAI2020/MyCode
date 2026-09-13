# -*- coding: utf-8 -*-
"""
创建科技感风格的PPT - 太原市网络空间安全协会第一届理事会工作总结
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 科技感配色
DARK_BLUE = RGBColor(10, 20, 50)
CYAN = RGBColor(0, 255, 255)
LIGHT_CYAN = RGBColor(100, 200, 255)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(180, 180, 180)

def set_dark_background(slide):
    """设置深色科技感背景"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE

def add_title(slide, text, top=Inches(0.5), font_size=44):
    """添加科技感标题"""
    title_box = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.333), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = text
    title_para.font.size = Pt(font_size)
    title_para.font.bold = True
    title_para.font.color.rgb = CYAN
    title_para.alignment = PP_ALIGN.CENTER

    # 添加底部发光线条
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3), top + Inches(0.7), Inches(7.333), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()

# ========== 第1页：封面 ==========
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide1)

# 装饰性网格线
for i in range(0, 13, 2):
    line = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(i), Inches(0), Inches(0.01), Inches(7.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 60, 120)
    line.line.fill.background()

# 主标题
title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(12.333), Inches(1.5))
title_frame = title_box.text_frame
title_para = title_frame.paragraphs[0]
title_para.text = "太原市网络空间安全协会"
title_para.font.size = Pt(54)
title_para.font.bold = True
title_para.font.color.rgb = CYAN
title_para.alignment = PP_ALIGN.CENTER

# 副标题
subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.333), Inches(1))
subtitle_frame = subtitle_box.text_frame
subtitle_para = subtitle_frame.paragraphs[0]
subtitle_para.text = "第一届理事会工作总结"
subtitle_para.font.size = Pt(40)
subtitle_para.font.bold = True
subtitle_para.font.color.rgb = WHITE
subtitle_para.alignment = PP_ALIGN.CENTER

# 年份
year_box = slide1.shapes.add_textbox(Inches(0.5), Inches(5.2), Inches(12.333), Inches(0.8))
year_frame = year_box.text_frame
year_para = year_frame.paragraphs[0]
year_para.text = "2024-2025"
year_para.font.size = Pt(36)
year_para.font.color.rgb = LIGHT_CYAN
year_para.alignment = PP_ALIGN.CENTER

# 底部信息
footer_box = slide1.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.333), Inches(0.6))
footer_frame = footer_box.text_frame
footer_para = footer_frame.paragraphs[0]
footer_para.text = "2025年12月"
footer_para.font.size = Pt(18)
footer_para.font.color.rgb = GRAY
footer_para.alignment = PP_ALIGN.CENTER

print("1. Cover page completed")

# ========== 第2页：目录 ==========
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide2)
add_title(slide2, "目录 CONTENTS")

toc_items = [
    ("01", "协会工作情况", "2024-2025年度工作回顾"),
    ("02", "工作亮点", "核心成就与突出贡献"),
    ("03", "存在问题", "当前面临的挑战"),
    ("04", "下一步工作计划", "未来发展蓝图"),
    ("05", "第二届理事会推荐名单", "新一届领导团队")
]

for i, (num, title, desc) in enumerate(toc_items):
    y_pos = Inches(1.5 + i * 1.1)

    # 序号圆圈
    circle = slide2.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(1), y_pos, Inches(0.8), Inches(0.8)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = CYAN
    circle.line.fill.background()

    num_box = slide2.shapes.add_textbox(Inches(1), y_pos + Inches(0.15), Inches(0.8), Inches(0.5))
    num_frame = num_box.text_frame
    num_para = num_frame.paragraphs[0]
    num_para.text = num
    num_para.font.size = Pt(24)
    num_para.font.bold = True
    num_para.font.color.rgb = DARK_BLUE
    num_para.alignment = PP_ALIGN.CENTER

    # 标题
    title_box = slide2.shapes.add_textbox(Inches(2.2), y_pos, Inches(5), Inches(0.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(24)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE

    # 描述
    desc_box = slide2.shapes.add_textbox(Inches(2.2), y_pos + Inches(0.45), Inches(8), Inches(0.4))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(14)
    desc_para.font.color.rgb = GRAY

print("2. Table of contents completed")

# ========== 第3页：2024年工作情况 ==========
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide3)
add_title(slide3, "2024年工作情况")

# 时间线项目
events_2024 = [
    ("7", "向太原市委网信办专题汇报", "明晰协会定位与发展路径"),
    ("9", "参与网络安全宣传周启动仪式", "普及网络安全知识"),
    ("9", "配合市公安局开展法制主题日", "强化市民网络安全法律意识"),
    ("10", "召开第一届会员大会", "正式成立协会"),
    ("10", "主办卫生健康网络安全培训班", "38家医疗机构140余人参与"),
    ("11", "赴市科协拜访", "提交加入市科协申请"),
    ("12", "参加企业社会责任蓝皮书发布会", "深化跨领域合作"),
    ("12", "联合举办国企等保测评培训", "市属国企150余人参加"),
]

# 左侧时间线
for i, (month, title, desc) in enumerate(events_2024[:4]):
    y_pos = Inches(1.4 + i * 1.4)

    # 圆点
    dot = slide3.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.8), y_pos + Inches(0.15), Inches(0.25), Inches(0.25)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = CYAN
    dot.line.fill.background()

    # 内容卡片
    card = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.3), y_pos, Inches(4.8), Inches(1.2)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0, 50, 100)
    card.fill.transparency = 0.3
    card.line.color.rgb = CYAN
    card.line.width = Pt(1)

    month_box = slide3.shapes.add_textbox(Inches(1.5), y_pos + Inches(0.1), Inches(4.4), Inches(0.35))
    month_frame = month_box.text_frame
    month_para = month_frame.paragraphs[0]
    month_para.text = f"Month {month}: {title}"
    month_para.font.size = Pt(14)
    month_para.font.bold = True
    month_para.font.color.rgb = CYAN

    desc_box = slide3.shapes.add_textbox(Inches(1.5), y_pos + Inches(0.5), Inches(4.4), Inches(0.5))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(11)
    desc_para.font.color.rgb = WHITE

# 右侧时间线
for i, (month, title, desc) in enumerate(events_2024[4:]):
    y_pos = Inches(1.4 + i * 1.4)

    # 圆点
    dot = slide3.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(6.8), y_pos + Inches(0.15), Inches(0.25), Inches(0.25)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = CYAN
    dot.line.fill.background()

    # 内容卡片
    card = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.3), y_pos, Inches(4.8), Inches(1.2)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0, 50, 100)
    card.fill.transparency = 0.3
    card.line.color.rgb = CYAN
    card.line.width = Pt(1)

    month_box = slide3.shapes.add_textbox(Inches(7.5), y_pos + Inches(0.1), Inches(4.4), Inches(0.35))
    month_frame = month_box.text_frame
    month_para = month_frame.paragraphs[0]
    month_para.text = f"Month {month}: {title}"
    month_para.font.size = Pt(14)
    month_para.font.bold = True
    month_para.font.color.rgb = CYAN

    desc_box = slide3.shapes.add_textbox(Inches(7.5), y_pos + Inches(0.5), Inches(4.4), Inches(0.5))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(11)
    desc_para.font.color.rgb = WHITE

print("3. 2024 Work completed")

# ========== 第4页：2025年工作情况 ==========
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide4)
add_title(slide4, "2025年协会成立后工作")

events_2025 = [
    ("1", "取得社会团体法人登记证书", "完成正式注册"),
    ("3", "向市委网信办和市公安局汇报", "获得主管部门支持"),
    ("4", "医疗机构安全巡检", "发现大量安全隐患"),
    ("5", "承办首席网络安全官培训", "市委网信办/市公安局/国资委主办"),
    ("5", "走访调研测评机构", "了解测评市场现状"),
    ("6", "成立党支部", "吴占分任党支部书记"),
    ("6", "签署等保测评自律公约", "省内六家机构"),
    ("7", "邀请云时代公司考察", "技术交流合作"),
    ("9", "国企入户安全辅导", "近20家市属国企"),
    ("10", "开通协会公众号", "增加宣传渠道"),
]

# 创建三列布局
cols = [(0.5, 3.5), (4.5, 3.5), (8.5, 3.5)]
col_events = [events_2025[:3], events_2025[3:6], events_2025[6:]]

for col_idx, (x_start, width) in enumerate(cols):
    events = col_events[col_idx]
    for i, (month, title, desc) in enumerate(events):
        y_pos = Inches(1.4 + i * 1.9)

        # 左边框线
        border = slide4.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x_start), y_pos, Inches(0.08), Inches(1.7)
        )
        border.fill.solid()
        border.fill.fore_color.rgb = CYAN
        border.line.fill.background()

        # 内容
        month_box = slide4.shapes.add_textbox(Inches(x_start + 0.3), y_pos, width - 0.5, Inches(0.4))
        month_frame = month_box.text_frame
        month_para = month_frame.paragraphs[0]
        month_para.text = f"Month {month}"
        month_para.font.size = Pt(16)
        month_para.font.bold = True
        month_para.font.color.rgb = CYAN

        title_box = slide4.shapes.add_textbox(Inches(x_start + 0.3), y_pos + Inches(0.4), width - 0.5, Inches(0.8))
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = title
        title_para.font.size = Pt(13)
        title_para.font.bold = True
        title_para.font.color.rgb = WHITE

        desc_box = slide4.shapes.add_textbox(Inches(x_start + 0.3), y_pos + Inches(1.1), width - 0.5, Inches(0.5))
        desc_frame = desc_box.text_frame
        desc_para = desc_frame.paragraphs[0]
        desc_para.text = desc
        desc_para.font.size = Pt(11)
        desc_para.font.color.rgb = GRAY

print("4. 2025 Work completed")

# ========== 第5页：工作亮点 ==========
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide5)
add_title(slide5, "工作亮点")

highlights = [
    ("01", "Efficient Setup", "From preparation to establishment\nOnly 5 months\nOutstanding coordination ability"),
    ("02", "Precise Services", "Focus on healthcare and enterprise sectors\nCovered nearly 500 professionals\nImproved industry protection level"),
    ("03", "Government Collaboration", "Deep cooperation with Cyberspace Office and Police\nParticipated in various activities\nEnhanced social credibility"),
]

for i, (num, title, desc) in enumerate(highlights):
    x_pos = Inches(1 + i * 4.2)

    # 背景卡片
    card = slide5.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x_pos, Inches(1.5), Inches(3.8), Inches(5)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0, 40, 80)
    card.fill.transparency = 0.4
    card.line.color.rgb = CYAN
    card.line.width = Pt(2)

    # 序号
    num_box = slide5.shapes.add_textbox(x_pos + Inches(0.3), Inches(1.8), Inches(3.2), Inches(0.8))
    num_frame = num_box.text_frame
    num_para = num_frame.paragraphs[0]
    num_para.text = num
    num_para.font.size = Pt(48)
    num_para.font.bold = True
    num_para.font.color.rgb = CYAN

    # 标题
    title_box = slide5.shapes.add_textbox(x_pos + Inches(0.3), Inches(2.8), Inches(3.2), Inches(0.6))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(26)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE
    title_para.alignment = PP_ALIGN.CENTER

    # 描述
    desc_box = slide5.shapes.add_textbox(x_pos + Inches(0.3), Inches(3.6), Inches(3.2), Inches(2.2))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(14)
    desc_para.font.color.rgb = GRAY
    desc_para.alignment = PP_ALIGN.CENTER

print("5. Highlights completed")

# ========== 第6页：存在问题 ==========
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide6)
add_title(slide6, "存在问题", font_size=40)

problems = [
    ("Incomplete Structure", "Lack of member base and visibility\nSelf-discipline committee not operational\nExpert committee not established\nParty branch activities missing"),
    ("Staff Shortage", "Lack of full-time staff\nDifficult to handle all work\nOften overwhelmed"),
    ("Funding Difficulties", "Salaries in arrears\nNo budget for accounting\nRent and registration paid by individuals"),
    ("Insufficient Activities", "Few academic exchanges and training\nLack of industry influence\nWeak public account promotion\nSlow progress on evaluation software"),
]

for i, (title, desc) in enumerate(problems):
    row = i // 2
    col = i % 2
    x_pos = Inches(0.5 + col * 6.5)
    y_pos = Inches(1.4 + row * 2.9)

    # 背景卡片
    card = slide6.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x_pos, y_pos, Inches(6), Inches(2.6)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(80, 20, 20)
    card.fill.transparency = 0.5
    card.line.color.rgb = RGBColor(255, 100, 100)
    card.line.width = Pt(2)

    # 标题
    title_box = slide6.shapes.add_textbox(x_pos + Inches(0.3), y_pos + Inches(0.2), Inches(5.4), Inches(0.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = f"WARNING: {title}"
    title_para.font.size = Pt(20)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(255, 150, 150)

    # 描述
    desc_box = slide6.shapes.add_textbox(x_pos + Inches(0.3), y_pos + Inches(0.8), Inches(5.4), Inches(1.6))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(13)
    desc_para.font.color.rgb = WHITE

print("6. Problems completed")

# ========== 第7页：下一步工作计划 ==========
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide7)
add_title(slide7, "下一步工作计划")

plans = [
    ("1", "Chairman Office Meeting", "Summarize work, propose restructuring, recommend candidates"),
    ("2", "Second Members Congress", "Elect new council members"),
    ("3", "Second Council First Meeting", "Elect chairman and vice chairmen, appoint secretary-general"),
    ("4", "Second Chairman Office Meeting", "Deploy next phase work"),
]

for i, (num, title, desc) in enumerate(plans):
    y_pos = Inches(1.4 + i * 1.45)

    # 序号
    circle = slide7.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.8), y_pos + Inches(0.1), Inches(0.7), Inches(0.7)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = CYAN
    circle.line.fill.background()

    num_box = slide7.shapes.add_textbox(Inches(0.8), y_pos + Inches(0.25), Inches(0.7), Inches(0.5))
    num_frame = num_box.text_frame
    num_para = num_frame.paragraphs[0]
    num_para.text = num
    num_para.font.size = Pt(24)
    num_para.font.bold = True
    num_para.font.color.rgb = DARK_BLUE
    num_para.alignment = PP_ALIGN.CENTER

    # 内容卡片
    card = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.8), y_pos, Inches(10.5), Inches(1.2)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0, 60, 120)
    card.fill.transparency = 0.3
    card.line.color.rgb = CYAN
    card.line.width = Pt(1.5)

    title_box = slide7.shapes.add_textbox(Inches(2.1), y_pos + Inches(0.15), Inches(10), Inches(0.45))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(20)
    title_para.font.bold = True
    title_para.font.color.rgb = CYAN

    desc_box = slide7.shapes.add_textbox(Inches(2.1), y_pos + Inches(0.65), Inches(10), Inches(0.45))
    desc_frame = desc_box.text_frame
    desc_para = desc_frame.paragraphs[0]
    desc_para.text = desc
    desc_para.font.size = Pt(14)
    desc_para.font.color.rgb = WHITE

print("7. Next steps completed")

# ========== 第8页：第二届理事会推荐名单 ==========
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide8)
add_title(slide8, "第二届理事会推荐名单", font_size=38)

# 左侧 - 领导层
leadership_title = slide8.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(6), Inches(0.5))
leadership_frame = leadership_title.text_frame
leadership_para = leadership_frame.paragraphs[0]
leadership_para.text = "Association Leadership"
leadership_para.font.size = Pt(22)
leadership_para.font.bold = True
leadership_para.font.color.rgb = CYAN

leaders = [
    ("Chairman", "Wang Fuming"),
    ("Honorary Chairman", "Chen Junjie"),
    ("Vice Chairmen", "Bai Shangwang, Liu Quanming, Jin Licong, Liu Yuntao, Yuan Xiaojun, Liu Dongxin, Zheng Yatong, Ren Lingyun"),
    ("Legal Representative", "Ren Lingyun"),
    ("Secretary General", "Qiu Jing"),
    ("Supervisory Chairman", "Wu Zhanfen"),
    ("Party Secretary", "Wu Zhanfen"),
]

y_start = Inches(1.9)
for i, (role, name) in enumerate(leaders):
    y_pos = y_start + i * 0.7
    role_box = slide8.shapes.add_textbox(Inches(0.5), y_pos, Inches(1.8), Inches(0.5))
    role_frame = role_box.text_frame
    role_para = role_frame.paragraphs[0]
    role_para.text = role + ":"
    role_para.font.size = Pt(14)
    role_para.font.bold = True
    role_para.font.color.rgb = LIGHT_CYAN

    name_box = slide8.shapes.add_textbox(Inches(2.3), y_pos, Inches(4), Inches(0.5))
    name_frame = name_box.text_frame
    name_para = name_frame.paragraphs[0]
    name_para.text = name
    name_para.font.size = Pt(13)
    name_para.font.color.rgb = WHITE

# 右侧 - 理事单位
org_title = slide8.shapes.add_textbox(Inches(6.8), Inches(1.3), Inches(6), Inches(0.5))
org_frame = org_title.text_frame
org_para = org_frame.paragraphs[0]
org_para.text = "Council Members (30 companies)"
org_para.font.size = Pt(22)
org_para.font.bold = True
org_para.font.color.rgb = CYAN

companies = [
    "Taiyuan Qingzhongxin Technology Co.",
    "Shanxi Jinxin'an Technology Co.",
    "Shanxi Haoyou Technology Development Co.",
    "Shanxi Saidun Network Security Evaluation Co.",
    "Shanxi Information Security Evaluation Center",
    "Shanxi Yinfumeixun Technology Co.",
    "Shanxi Lanying Technology Co.",
    "Shanxi Sanyouhe Smart Information Technology Co.",
    "Shanxi Youxin Network Security Technology Co.",
    "Shanxi Zhongchengda Information Service Co.",
    "...(Total 30 companies)",
]

y_start = Inches(1.9)
for i, company in enumerate(companies):
    y_pos = y_start + i * 0.52
    if y_pos > Inches(7):
        break
    company_box = slide8.shapes.add_textbox(Inches(6.8), y_pos, Inches(6), Inches(0.5))
    company_frame = company_box.text_frame
    company_para = company_frame.paragraphs[0]
    company_para.text = f"* {company}"
    company_para.font.size = Pt(12)
    company_para.font.color.rgb = WHITE

print("8. Council members completed")

# ========== 第9页：结束页 ==========
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
set_dark_background(slide9)

# 装饰性网格
for i in range(0, 13, 2):
    line = slide9.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(i), Inches(0), Inches(0.01), Inches(7.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 60, 120)
    line.line.fill.background()

# 感谢文字
thanks_box = slide9.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1))
thanks_frame = thanks_box.text_frame
thanks_para = thanks_frame.paragraphs[0]
thanks_para.text = "THANK YOU"
thanks_para.font.size = Pt(60)
thanks_para.font.bold = True
thanks_para.font.color.rgb = CYAN
thanks_para.alignment = PP_ALIGN.CENTER

# 副标题
sub_box = slide9.shapes.add_textbox(Inches(0.5), Inches(4), Inches(12.333), Inches(0.8))
sub_frame = sub_box.text_frame
sub_para = sub_frame.paragraphs[0]
sub_para.text = "Taiyuan Cyberspace Security Association"
sub_para.font.size = Pt(32)
sub_para.font.color.rgb = WHITE
sub_para.alignment = PP_ALIGN.CENTER

# 底部
footer_box = slide9.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(0.6))
footer_frame = footer_box.text_frame
footer_para = footer_frame.paragraphs[0]
footer_para.text = "Building Cybersecurity Ecosystem Together"
footer_para.font.size = Pt(20)
footer_para.font.color.rgb = GRAY
footer_para.alignment = PP_ALIGN.CENTER

print("9. Ending page completed")

# 保存文件
output_path = r"d:\python files\Association_Work_Summary.pptx"
prs.save(output_path)
print(f"\nPPT saved to: {output_path}")
print(f"File size: {os.path.getsize(output_path) / 1024:.1f} KB")
