# -*- coding: utf-8 -*-
"""报告格式转换器 - 支持 HTML/PDF/Word/MarkDown"""
import os
import re
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def convert_html_to_markdown(html: str) -> str:
    """HTML → MarkDown 转换"""
    md = html

    # 提取title
    title_match = re.search(r'<title>(.*?)</title>', md, re.DOTALL)
    title = title_match.group(1).strip() if title_match else '报告'

    # 移除 <head> 和 <style> 块
    md = re.sub(r'<head>.*?</head>', '', md, flags=re.DOTALL)
    md = re.sub(r'<style>.*?</style>', '', md, flags=re.DOTALL)
    md = re.sub(r'<script>.*?</script>', '', md, flags=re.DOTALL)

    # 标题转换
    md = re.sub(r'<h1[^>]*>(.*?)</h1>', r'# \1\n', md, flags=re.DOTALL)
    md = re.sub(r'<h2[^>]*>(.*?)</h2>', r'## \1\n', md, flags=re.DOTALL)
    md = re.sub(r'<h3[^>]*>(.*?)</h3>', r'### \1\n', md, flags=re.DOTALL)
    md = re.sub(r'<h4[^>]*>(.*?)</h4>', r'#### \1\n', md, flags=re.DOTALL)

    # 段落
    md = re.sub(r'<p[^>]*>(.*?)</p>', r'\1\n\n', md, flags=re.DOTALL)

    # 换行
    md = re.sub(r'<br\s*/?>', '\n', md)

    # 先处理内联元素(在表格转换之前)
    # 链接
    md = re.sub(r'<a[^>]*href=["\'](.*?)["\'][^>]*>(.*?)</a>', r'[\2](\1)', md, flags=re.DOTALL)
    # 粗体/斜体
    md = re.sub(r'<(b|strong)[^>]*>(.*?)</\1>', r'**\2**', md, flags=re.DOTALL)
    md = re.sub(r'<(i|em)[^>]*>(.*?)</\1>', r'*\2*', md, flags=re.DOTALL)
    # 代码
    md = re.sub(r'<code[^>]*>(.*?)</code>', r'`\1`', md, flags=re.DOTALL)

    # 表格 → Markdown 表格
    def table_to_md(match):
        table_html = match.group(0)
        rows = re.findall(r'<tr[^>]*>(.*?)</tr>', table_html, re.DOTALL)
        if not rows:
            return '\n'
        md_rows = []
        for i, row in enumerate(rows):
            cells = re.findall(r'<t[hd][^>]*>(.*?)</t[hd]>', row, re.DOTALL)
            cells = [re.sub(r'<[^>]+>', '', c).strip() for c in cells]
            md_rows.append('| ' + ' | '.join(cells) + ' |')
            if i == 0:
                md_rows.append('|' + '|'.join([' --- ' for _ in cells]) + '|')
        return '\n'.join(md_rows) + '\n'

    md = re.sub(r'<table[^>]*>.*?</table>', table_to_md, md, flags=re.DOTALL)

    # 列表
    md = re.sub(r'<li[^>]*>(.*?)</li>', r'- \1\n', md, flags=re.DOTALL)

    # 删除剩余HTML标签
    md = re.sub(r'<[^>]+>', '', md)

    # 清理多余空行
    md = re.sub(r'\n{3,}', '\n\n', md)

    # 添加标题
    if not md.strip().startswith('# '):
        md = f'# {title}\n\n' + md

    return md.strip()


def convert_html_to_pdf(html: str, output_path: str) -> bool:
    """HTML → PDF 转换 (使用 weasyprint)"""
    try:
        from weasyprint import HTML
        HTML(string=html).write_pdf(output_path)
        return True
    except (ImportError, OSError) as e:
        # weasyprint 需要 GTK+ 系统库，Windows 下可能不可用
        # 回退：保存为 HTML 格式
        fallback = output_path.rsplit('.', 1)[0] + '.html'
        with open(fallback, 'w', encoding='utf-8') as f:
            f.write(html)
        logger.warning(f"PDF生成需要GTK+库(weasyprint依赖)，已保存为HTML: {fallback}")
        return False
    except Exception as e:
        logger.error(f"PDF生成失败: {e}")
        # 回退保存为HTML
        fallback = output_path.rsplit('.', 1)[0] + '.html'
        with open(fallback, 'w', encoding='utf-8') as f:
            f.write(html)
        return False


def convert_html_to_docx(html: str, output_path: str) -> bool:
    """HTML → Word (.docx) 转换 (使用 python-docx)"""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import html as html_mod

        doc = Document()

        # 提取body内容
        body_match = re.search(r'<body[^>]*>(.*?)</body>', html, re.DOTALL)
        content = body_match.group(1) if body_match else html

        # 移除style/script
        content = re.sub(r'<style[^>]*>.*?</style>', '', content, flags=re.DOTALL)
        content = re.sub(r'<script[^>]*>.*?</script>', '', content, flags=re.DOTALL)

        # 分段处理
        # 先替换标题
        content = re.sub(r'<h1[^>]*>(.*?)</h1>', r'[H1]\1[/H1]', content, flags=re.DOTALL)
        content = re.sub(r'<h2[^>]*>(.*?)</h2>', r'[H2]\1[/H2]', content, flags=re.DOTALL)
        content = re.sub(r'<h3[^>]*>(.*?)</h3>', r'[H3]\1[/H3]', content, flags=re.DOTALL)
        # 段落
        content = re.sub(r'<p[^>]*>(.*?)</p>', r'\1[PARA]', content, flags=re.DOTALL)
        # 换行
        content = re.sub(r'<br[^>]*>', '\n', content)
        # 列表项
        content = re.sub(r'<li[^>]*>(.*?)</li>', r'• \1\n', content, flags=re.DOTALL)
        # div结束也加换行
        content = re.sub(r'</div>', '\n', content)
        # 表格简化
        content = re.sub(r'<table[^>]*>.*?</table>', '[表格内容]', content, flags=re.DOTALL)
        # 移除其他HTML标签
        content = re.sub(r'<[^>]+>', '', content)
        content = html_mod.unescape(content)

        # 按段落分割
        paragraphs = [p.strip() for p in content.split('[PARA]') if p.strip()]

        for para_text in paragraphs:
            para_text = para_text.strip()
            if not para_text:
                continue

            if para_text.startswith('[H1]') and '[/H1]' in para_text:
                text = para_text.replace('[H1]', '').replace('[/H1]', '').strip()
                p = doc.add_heading(text, level=1)
            elif para_text.startswith('[H2]') and '[/H2]' in para_text:
                text = para_text.replace('[H2]', '').replace('[/H2]', '').strip()
                p = doc.add_heading(text, level=2)
            elif para_text.startswith('[H3]') and '[/H3]' in para_text:
                text = para_text.replace('[H3]', '').replace('[/H3]', '').strip()
                p = doc.add_heading(text, level=3)
            else:
                p = doc.add_paragraph(para_text)
                p.style.font.size = Pt(11)

        doc.save(output_path)
        return True
    except ImportError:
        logger.warning("python-docx未安装，无法生成Word文档")
        return False
    except Exception as e:
        logger.error(f"Word生成失败: {e}")
        return False


def save_report_formatted(html_content: str, filepath: str) -> tuple:
    """根据文件扩展名保存为对应格式。返回 (success, actual_path)"""
    ext = os.path.splitext(filepath)[1].lower()

    if ext in ('.html', '.htm'):
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(html_content)
        return (True, filepath)

    elif ext == '.md':
        md_content = convert_html_to_markdown(html_content)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(md_content)
        return (True, filepath)

    elif ext == '.pdf':
        ok = convert_html_to_pdf(html_content, filepath)
        if ok:
            return (True, filepath)
        else:
            # PDF失败，回退到HTML
            fallback = filepath.rsplit('.', 1)[0] + '.html'
            return (True, fallback)

    elif ext == '.docx':
        ok = convert_html_to_docx(html_content, filepath)
        if ok:
            return (True, filepath)
        else:
            fallback = filepath.rsplit('.', 1)[0] + '.html'
            with open(fallback, 'w', encoding='utf-8') as f:
                f.write(html_content)
            return (True, fallback)

    else:
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(html_content)
        return (True, filepath)
