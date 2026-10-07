"""报告台账 CRUD + 服务报告文件（脱敏 + 加密入库）。"""
import hashlib
import html as html_mod
import io
import json
import os
import re
from urllib.parse import quote

from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile, status
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.deps import assert_scoped, customer_scope_of, get_db, require_role, scope_filter
from app.core.security import mask_text
from app.models import CmdbCi, Contract, ContractItem, Customer, Issue, Report, SysUser, WorkOrder, WorkOrderCi, WorkOrderCycle, WorkOrderItem
from app.schemas.report import ReportCreate, ReportOut, ReportUpdate
from app.services.archive_crypto import decrypt_bytes, encrypt_bytes
from app.services.audit_service import record
from app.services.contract_extract_service import extract_text
from app.utils.response import ok

ROLE = ("sys_admin", "sys_ops", "ticket_mgr")
READ_ROLE = ("sys_admin", "sys_ops", "ticket_mgr", "cust_admin", "cust_service")

router = APIRouter(prefix="/reports", tags=["报告中心"])

# 允许上传的扩展名 → 派生 MIME（不信任客户端 content_type）
REPORT_MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
}

# 仅可信二进制/位图可内联渲染；md/html 强制下载（防存储型 XSS）
ALLOWED_INLINE_MIME = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "image/jpeg",
    "image/png",
}

REPORTS_DIR = os.path.join(settings.UPLOAD_DIR, "reports")


def _fmt_date(d) -> str:
    return f"{d.year}年{d.month}月{d.day}日" if d else ""


def _report_context(db: Session, wo: WorkOrder) -> dict:
    """聚合生成报告所需的上下文：客户/项目/业务系统/服务类别/起止日期。"""
    customer_name = ""
    if wo.customer_id:
        _c = db.get(Customer, wo.customer_id)
        customer_name = _c.name if _c else ""

    project_name = ""
    contract_id = wo.contract_id
    if contract_id is None:
        _it = db.query(WorkOrderItem.contract_item_id).filter_by(work_order_id=wo.id).first()
        if _it:
            _ci = db.get(ContractItem, _it[0])
            if _ci:
                contract_id = _ci.contract_id
    if contract_id:
        _c = db.get(Contract, contract_id)
        project_name = _c.name if _c else ""

    ci_names: list[str] = []
    for r in db.query(WorkOrderCi).filter_by(work_order_id=wo.id).all():
        _ci = db.get(CmdbCi, r.ci_id)
        if _ci:
            ci_names.append(_ci.name)
    if not ci_names and wo.ci_id is not None:
        _ci = db.get(CmdbCi, wo.ci_id)
        if _ci:
            ci_names.append(_ci.name)

    service_names: list[str] = []
    seen: set[str] = set()
    for r in db.query(WorkOrderItem).filter_by(work_order_id=wo.id).all():
        _item = db.get(ContractItem, r.contract_item_id)
        if _item and _item.project and _item.project not in seen:
            seen.add(_item.project)
            service_names.append(_item.project)

    start_date = end_date = None
    for r in db.query(WorkOrderCycle).filter_by(work_order_id=wo.id).all():
        if r.service_start and (start_date is None or r.service_start < start_date):
            start_date = r.service_start
        if r.service_end and (end_date is None or r.service_end > end_date):
            end_date = r.service_end

    return {
        "customer_name": customer_name,
        "project_name": project_name,
        "ci_names": ci_names,
        "service_names": service_names,
        "start_date": start_date,
        "end_date": end_date,
    }


def _build_title(ctx: dict) -> str:
    parts = []
    if ctx["customer_name"]:
        parts.append(ctx["customer_name"])
    if ctx["project_name"]:
        parts.append(ctx["project_name"])
    if ctx["ci_names"]:
        parts.append("、".join(ctx["ci_names"]))
    if ctx["start_date"] and ctx["end_date"]:
        parts.append(f"{_fmt_date(ctx['start_date'])}-{_fmt_date(ctx['end_date'])}")
    if ctx["service_names"]:
        parts.append("、".join(ctx["service_names"]))
    return "".join(parts) + "工作报告"


def _generate_html_report(ctx: dict, rd: dict) -> str:
    def esc(s) -> str:
        return html_mod.escape(str(s if s is not None else ""))

    levels = ["严重", "高危", "中危", "低危", "其他"]
    cats = ["漏洞", "配置缺陷", "风险隐患", "基线不合规"]
    issues = rd.get("issues") or {}
    rows = []
    for cat in cats:
        lv = issues.get(cat) or {}
        cells = "".join(f"<td>{esc(lv.get(l, 0))}</td>" for l in levels)
        rows.append(f"<tr><th>{esc(cat)}</th>{cells}</tr>")
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{esc(ctx['title'])}</title>
<style>
body{{font-family:"Microsoft YaHei","PingFang SC",sans-serif;color:#222;margin:40px;line-height:1.7}}
h1{{font-size:24px;border-bottom:2px solid #333;padding-bottom:8px}}
h2{{font-size:18px;margin-top:24px}}
table{{border-collapse:collapse;width:100%;margin-top:8px}}
th,td{{border:1px solid #999;padding:6px 10px;text-align:center}}
th{{background:#f2f2f2}}
.meta{{color:#666;font-size:13px}}
</style></head><body>
<h1>{esc(ctx['title'])}</h1>
<p class="meta">客户：{esc(ctx['customer_name'])}　项目：{esc(ctx['project_name'])}　业务系统：{esc('、'.join(ctx['ci_names']))}　服务类别：{esc('、'.join(ctx['service_names']))}</p>
<p class="meta">服务周期：{esc(_fmt_date(ctx['start_date']))} 至 {esc(_fmt_date(ctx['end_date']))}</p>
<h2>一、工作内容</h2>
<p>{esc(rd.get('work_content') or '—')}</p>
<h2>二、安全问题统计</h2>
<table><thead><tr><th>类别</th>{"".join(f"<th>{esc(l)}</th>" for l in levels)}</tr></thead><tbody>{rows}</tbody></table>
<h2>三、详细情况</h2>
<p>{esc(rd.get('details') or '—')}</p>
</body></html>"""


def _generate_docx_report(ctx: dict, rd: dict) -> bytes:
    doc = Document()
    doc.add_heading(ctx["title"], level=0)
    doc.add_paragraph(
        f"客户：{ctx['customer_name']}　项目：{ctx['project_name']}　业务系统：{'、'.join(ctx['ci_names'])}　服务类别：{'、'.join(ctx['service_names'])}"
    )
    doc.add_paragraph(f"服务周期：{_fmt_date(ctx['start_date'])} 至 {_fmt_date(ctx['end_date'])}")

    doc.add_heading("一、工作内容", level=1)
    doc.add_paragraph(rd.get("work_content") or "—")

    doc.add_heading("二、安全问题统计", level=1)
    levels = ["严重", "高危", "中危", "低危", "其他"]
    cats = ["漏洞", "配置缺陷", "风险隐患", "基线不合规"]
    issues = rd.get("issues") or {}
    table = doc.add_table(rows=1, cols=len(levels) + 1)
    hdr = table.rows[0].cells
    hdr[0].text = "类别"
    for i, l in enumerate(levels):
        hdr[i + 1].text = l
    for cat in cats:
        lv = issues.get(cat) or {}
        row = table.add_row().cells
        row[0].text = cat
        for i, l in enumerate(levels):
            row[i + 1].text = str(lv.get(l, 0))

    doc.add_heading("三、详细情况", level=1)
    doc.add_paragraph(rd.get("details") or "—")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _html_to_text(s: str) -> str:
    return html_mod.unescape(re.sub(r"<[^>]+>", " ", s))


def _extract_report_text(filename: str, data: bytes) -> str:
    """按扩展名抽取报告纯文本：md/html 直接解码，pdf/docx/图片复用合同抽取。"""
    lower = (filename or "").lower()
    if lower.endswith((".md", ".html", ".htm")):
        for enc in ("utf-8", "utf-8-sig", "gbk"):
            try:
                text = data.decode(enc)
                return _html_to_text(text) if lower.endswith((".html", ".htm")) else text
            except UnicodeDecodeError:
                continue
        return ""
    try:
        text, _ = extract_text(filename, data)
        return text
    except ValueError:
        return ""


def _enrich(db: Session, items: list[dict]) -> None:
    cids = {it["customer_id"] for it in items if it["customer_id"]}
    pids = {it["contract_id"] for it in items if it["contract_id"]}
    wids = {it["work_order_id"] for it in items if it["work_order_id"]}
    cname = {c.id: c.name for c in db.query(Customer).filter(Customer.id.in_(cids)).all()} if cids else {}
    pname = {c.id: c.name for c in db.query(Contract).filter(Contract.id.in_(pids)).all()} if pids else {}
    wno = {w.id: w.no for w in db.query(WorkOrder).filter(WorkOrder.id.in_(wids)).all()} if wids else {}
    for it in items:
        it["customer_name"] = cname.get(it["customer_id"])
        it["project_name"] = pname.get(it["contract_id"])
        it["work_order_no"] = wno.get(it["work_order_id"])


@router.get("")
def list_reports(
    report_type: str | None = Query(None),
    status: str | None = Query(None),
    work_order_id: int | None = Query(None),
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    user: SysUser = Depends(require_role(*READ_ROLE)),
    db: Session = Depends(get_db),
):
    scope = customer_scope_of(user, db)
    q = scope_filter(db.query(Report), Report, scope)
    if report_type:
        q = q.filter(Report.report_type == report_type)
    if status:
        q = q.filter(Report.status == status)
    if work_order_id is not None:
        q = q.filter(Report.work_order_id == work_order_id)
    total = q.count()
    rows = q.order_by(Report.id.desc()).offset((page - 1) * size).limit(size).all()
    items = []
    for r in rows:
        d = ReportOut.model_validate(r).model_dump()
        d["has_file"] = r.content_enc is not None
        items.append(d)
    data = {"items": items, "total": total, "page": page, "size": size}
    _enrich(db, data["items"])
    return ok(data)


@router.post("/upload")
def upload_report(
    file: UploadFile = File(...),
    work_order_id: int | None = Form(None),
    report_type: str = Form("运维报告"),
    title: str | None = Form(None),
    report_data: str | None = Form(None),
    user: SysUser = Depends(require_role(*ROLE)),
    db: Session = Depends(get_db),
):
    """上传服务报告：提取文本 → 脱敏 → Fernet 加密入库，挂在工单（反推项目/客户）下。"""
    filename = (file.filename or "").strip()
    ext = os.path.splitext(filename.lower())[1]
    if ext not in REPORT_MIME_BY_EXT:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "仅支持 PDF / Word(.docx) / Markdown(.md) / HTML / 图片(jpg/png) 文件")
    data = file.file.read()
    if not data:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "文件为空")
    max_bytes = settings.ARCHIVE_MAX_MB * 1024 * 1024
    if len(data) > max_bytes:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"文件不能超过 {settings.ARCHIVE_MAX_MB}MB")

    mime = REPORT_MIME_BY_EXT[ext]
    raw_text = _extract_report_text(filename, data)
    masked = mask_text(raw_text)

    contract_id = customer_id = None
    wo = None
    ctx = None
    if work_order_id is not None:
        wo = db.get(WorkOrder, work_order_id)
        if wo is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "工单不存在")
        assert_scoped(wo, customer_scope_of(user, db), db)
        contract_id = wo.contract_id
        customer_id = wo.customer_id
        if customer_id is None and contract_id is not None:
            _c = db.get(Contract, contract_id)
            customer_id = _c.customer_id if _c else None
        ctx = _report_context(db, wo)
        ctx["title"] = _build_title(ctx)

    # 结构化报告内容同样脱敏：工作内容/详细情况含敏感信息（漏洞 IP、联系方式等），不落明文
    if report_data:
        try:
            rd = json.loads(report_data)
            if isinstance(rd, dict):
                for k in ("work_content", "details"):
                    if isinstance(rd.get(k), str):
                        rd[k] = mask_text(rd[k])
                report_data = json.dumps(rd, ensure_ascii=False)
        except (ValueError, TypeError):
            pass

    generated_title = ctx["title"] if ctx else ""
    obj = Report(
        title=generated_title or (title or "").strip() or os.path.splitext(filename)[0],
        report_type=report_type,
        customer_id=customer_id,
        contract_id=contract_id,
        work_order_id=work_order_id,
        status="已提交",
        summary=(masked[:500] if masked else None),
        original_filename=filename,
        mime_type=mime,
        file_size=len(data),
        file_hash=hashlib.sha256(data).hexdigest(),
        content_enc=encrypt_bytes(data),
        masked_text=masked,
        report_data=report_data,
    )
    db.add(obj)
    db.flush()

    # 生成 HTML + Word 报告（内容已脱敏）并加密落盘 /reports
    if ctx:
        rd = json.loads(report_data) if report_data else {}
        html_bytes = _generate_html_report(ctx, rd).encode("utf-8")
        docx_bytes = _generate_docx_report(ctx, rd)
        os.makedirs(REPORTS_DIR, exist_ok=True)
        html_name = f"{obj.id}_report.html.enc"
        docx_name = f"{obj.id}_report.docx.enc"
        with open(os.path.join(REPORTS_DIR, html_name), "wb") as f:
            f.write(encrypt_bytes(html_bytes))
        with open(os.path.join(REPORTS_DIR, docx_name), "wb") as f:
            f.write(encrypt_bytes(docx_bytes))
        obj.html_filename = html_name
        obj.docx_filename = docx_name

    # 安全问题统计 → 风险管控问题（按类型聚合，级别数量存 level_counts）
    if work_order_id is not None and report_data:
        try:
            rd = json.loads(report_data)
            issues_map = rd.get("issues") or {}
            if isinstance(issues_map, dict):
                type_map = {"漏洞": "安全漏洞"}
                for cat, lv in issues_map.items():
                    if not isinstance(lv, dict):
                        continue
                    counts = {k: int(v) for k, v in lv.items() if isinstance(v, (int, float)) and v}
                    if not counts:
                        continue
                    db.add(Issue(
                        work_order_id=work_order_id,
                        type=type_map.get(cat, cat),
                        level=max(counts, key=counts.get),
                        description=rd.get("details") or f"{cat}隐患统计",
                        level_counts=json.dumps(counts, ensure_ascii=False),
                        status="待整改",
                    ))
        except (ValueError, TypeError):
            pass

    record(db, user_id=user.id, action="upload_report", resource=f"report:{obj.id}", after=filename)
    db.commit()
    return ok(ReportOut.model_validate(obj).model_dump())


@router.post("")
def create_report(
    body: ReportCreate,
    user: SysUser = Depends(require_role(*ROLE)),
    db: Session = Depends(get_db),
):
    obj = Report(**body.model_dump())
    db.add(obj)
    db.flush()
    record(db, user_id=user.id, action="create", resource=f"report:{obj.id}", after=str(body.model_dump()))
    db.commit()
    return ok(ReportOut.model_validate(obj).model_dump())


@router.get("/{rid}/preview")
def preview_report(rid: int, user: SysUser = Depends(require_role(*READ_ROLE)), db: Session = Depends(get_db)):
    """返回报告脱敏文本，供在线预览/搜索（不泄露敏感信息）。"""
    obj = db.get(Report, rid)
    if obj is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在")
    scope = customer_scope_of(user, db)
    if scope is not None and obj.customer_id != scope:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在或无权访问")
    return ok({
        "title": obj.title,
        "original_filename": obj.original_filename,
        "mime_type": obj.mime_type,
        "masked_text": obj.masked_text or "",
    })


@router.get("/{rid}/download")
def download_report(rid: int, user: SysUser = Depends(require_role(*READ_ROLE)), db: Session = Depends(get_db)):
    """解密下载报告原文。"""
    obj = db.get(Report, rid)
    if obj is None or obj.content_enc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告或附件不存在")
    scope = customer_scope_of(user, db)
    if scope is not None and obj.customer_id != scope:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在或无权访问")
    data = decrypt_bytes(obj.content_enc)
    mime = obj.mime_type if obj.mime_type in ALLOWED_INLINE_MIME else "application/octet-stream"
    disposition = "inline" if mime != "application/octet-stream" else "attachment"
    return Response(
        content=data,
        media_type=mime,
        headers={
            "Content-Disposition": f"{disposition}; filename*=UTF-8''{quote(obj.original_filename or 'report')}",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.get("/{rid}/generated/{kind}")
def download_generated_report(rid: int, kind: str, user: SysUser = Depends(require_role(*READ_ROLE)), db: Session = Depends(get_db)):
    """下载生成的 HTML/Word 报告（解密）。kind: html | docx"""
    obj = db.get(Report, rid)
    if obj is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在")
    scope = customer_scope_of(user, db)
    if scope is not None and obj.customer_id != scope:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在或无权访问")
    if kind == "html":
        fname, mime, fallback = obj.html_filename, "text/html", "report.html"
    elif kind == "docx":
        fname, mime, fallback = obj.docx_filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "report.docx"
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "kind 仅支持 html / docx")
    if not fname:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告文档不存在")
    path = os.path.join(REPORTS_DIR, fname)
    if not os.path.exists(path):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告文档不存在")
    with open(path, "rb") as f:
        data = decrypt_bytes(f.read())
    return Response(
        content=data,
        media_type=mime,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fallback)}", "X-Content-Type-Options": "nosniff"},
    )


@router.put("/{rid}")
def update_report(
    rid: int,
    body: ReportUpdate,
    user: SysUser = Depends(require_role(*ROLE)),
    db: Session = Depends(get_db),
):
    obj = db.get(Report, rid)
    if obj is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在")
    data = body.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(obj, k, v)
    db.flush()
    record(db, user_id=user.id, action="update", resource=f"report:{rid}", after=str(data))
    db.commit()
    return ok(ReportOut.model_validate(obj).model_dump())


@router.delete("/{rid}")
def delete_report(
    rid: int,
    user: SysUser = Depends(require_role("sys_admin", "sys_ops")),
    db: Session = Depends(get_db),
):
    obj = db.get(Report, rid)
    if obj is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "报告不存在")
    db.delete(obj)
    record(db, user_id=user.id, action="delete", resource=f"report:{rid}")
    db.commit()
    return ok({"deleted": rid})
